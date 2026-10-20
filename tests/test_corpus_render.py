from pathlib import Path

import docx
from pypdf import PdfReader

from rootline.core.models import SopRecord
from rootline.corpus.docx_writer import write_docx
from rootline.corpus.layout import SOP_SECTIONS, Heading, Layout, layout_for, load_source
from rootline.corpus.md_writer import to_markdown, write_markdown
from rootline.corpus.pdf_writer import write_pdf

CORPUS = Path(__file__).parents[1] / "corpus"
SOURCE_DIR = CORPUS / "sources" / "greenfield"
OUTPUT_DIR = CORPUS / "greenfield"
SOP_SOURCE = SOURCE_DIR / "sop-db-001.yaml"
RCA_SOURCE = SOURCE_DIR / "rca-2025-014.yaml"


def sop_steps() -> list[str]:
    record = load_source(SOP_SOURCE).record
    assert isinstance(record, SopRecord)
    return [step.instruction for step in record.steps]


def headings(layout: Layout, level: int) -> list[str]:
    return [e.text for e in layout.elements if isinstance(e, Heading) and e.level == level]


def appear_in_order(text: str, parts: list[str]) -> bool:
    flat = " ".join(text.split())
    positions = [flat.find(" ".join(part.split())) for part in parts]
    return -1 not in positions and positions == sorted(positions)


def test_sop_layout_has_every_template_section_in_order() -> None:
    numbered = [h for h in headings(layout_for(load_source(SOP_SOURCE)), 2) if h[0].isdigit()]

    assert numbered == [f"{n}. {name}" for n, name in enumerate(SOP_SECTIONS, start=1)]


def test_rca_layout_follows_template_order() -> None:
    assert headings(layout_for(load_source(RCA_SOURCE)), 2) == [
        "Summary",
        "Problem statement",
        "Impact",
        "Immediate resolution",
        "Event timeline",
        "Analysis",
        "Root causes",
        "Action items",
    ]


def test_markdown_keeps_procedure_steps_in_order(tmp_path: Path) -> None:
    target = tmp_path / "sop.md"

    write_markdown(layout_for(load_source(SOP_SOURCE)), target)

    assert appear_in_order(target.read_text(), sop_steps())


def test_docx_procedure_table_matches_steps(tmp_path: Path) -> None:
    target = tmp_path / "sop.docx"

    write_docx(layout_for(load_source(SOP_SOURCE)), target)

    document = docx.Document(str(target))
    procedure = next(t for t in document.tables if t.rows[0].cells[0].text == "Step")
    assert [row.cells[1].text for row in procedure.rows[1:]] == sop_steps()


def test_docx_page_header_carries_document_number(tmp_path: Path) -> None:
    target = tmp_path / "sop.docx"

    write_docx(layout_for(load_source(SOP_SOURCE)), target)

    header = docx.Document(str(target)).sections[0].header
    assert any("SOP-DB-001" in paragraph.text for paragraph in header.paragraphs)


def test_pdf_text_keeps_procedure_steps_in_order(tmp_path: Path) -> None:
    target = tmp_path / "sop.pdf"

    write_pdf(layout_for(load_source(SOP_SOURCE)), target)

    text = "\n".join(page.extract_text() for page in PdfReader(target).pages)
    assert appear_in_order(text, sop_steps())


def test_pdf_footer_counts_total_pages(tmp_path: Path) -> None:
    target = tmp_path / "sop.pdf"

    write_pdf(layout_for(load_source(SOP_SOURCE)), target)

    pages = PdfReader(target).pages
    assert f"Page 1 of {len(pages)}" in pages[0].extract_text()


def test_every_source_has_exactly_one_rendered_file() -> None:
    expected = {f"{path.stem}.{load_source(path).render_as}" for path in SOURCE_DIR.glob("*.yaml")}

    assert {path.name for path in OUTPUT_DIR.iterdir()} == expected


def test_committed_markdown_matches_sources() -> None:
    for path in SOURCE_DIR.glob("*.yaml"):
        source = load_source(path)
        if source.render_as == "md":
            committed = (OUTPUT_DIR / f"{path.stem}.md").read_text()
            assert committed == to_markdown(layout_for(source)), path.name

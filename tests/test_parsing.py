from pathlib import Path

import docx
import pytest

from rootline.corpus.layout import load_document
from rootline.ingest.hashing import content_hash, raw_sha256, store_raw
from rootline.ingest.parsing import UnsupportedFormat, parse_bytes, parse_file
from rootline.models import ParsedDoc, SopRecord

CORPUS = Path(__file__).parents[1] / "corpus"
CORPUS_FILES = sorted(
    path
    for folder in ("greenfield", "brownfield", "reference")
    for path in (CORPUS / folder).iterdir()
)


def parse(relative: str) -> ParsedDoc:
    return parse_file(CORPUS / relative)


def source_steps(doc_id: str) -> list[str]:
    _, record = load_document(CORPUS / "sources" / "greenfield" / f"{doc_id}.yaml")
    assert isinstance(record, SopRecord)
    return [step.instruction for step in record.steps]


def procedure_rows(doc: ParsedDoc) -> list[list[str]]:
    table = next(b for b in doc.blocks if b.kind == "table" and b.rows[0][0] == "Step")
    return table.rows[1:]


@pytest.mark.parametrize("path", CORPUS_FILES, ids=lambda path: path.name)
def test_every_corpus_file_parses_with_matching_offsets(path: Path) -> None:
    doc = parse_file(path)

    assert doc.title and doc.blocks
    for block in doc.blocks:
        assert doc.text[block.start : block.end] == block.text


def test_markdown_front_matter_becomes_fields() -> None:
    doc = parse("reference/ref-k8s-001.md")

    assert doc.fields["id"] == "REF-K8S-001"
    assert doc.title == "Node-pressure eviction when a node runs low on disk"


def test_markdown_table_cells_lose_code_backticks() -> None:
    commands = [row[2] for row in procedure_rows(parse("greenfield/sop-cache-001.md"))]

    assert commands[0] == "redis-cli -h cache ping"


def test_text_runbook_steps_keep_their_commands() -> None:
    items = [b.text for b in parse("brownfield/orders-db-connections-runbook.txt").blocks]
    steps = [text for text in items if text[:2] in {"1)", "2)", "3)", "4)", "5)"}]

    assert len(steps) == 5
    assert "DB_POOL_MAX=50" in steps[3]


def test_html_drops_wiki_boilerplate_and_keeps_last_modified() -> None:
    doc = parse("brownfield/wiki-orders-api-db-pool.html")

    assert "Log in" not in doc.text
    assert "Exported from the wiki" not in doc.text
    assert doc.fields["last_modified"] == "2024-03-02"


def test_email_title_drops_reply_prefix_and_body_is_decoded() -> None:
    doc = parse("brownfield/re-orders-api-503s-this-morning.eml")

    assert doc.title == "orders api outage from exhausted database connection pool"
    assert "What happened: For 38 minutes orders-api could not get database connections" in doc.text


def test_slack_messages_become_timestamped_paragraphs() -> None:
    doc = parse("brownfield/inc-2025-06-03-orders-db-pool.json")
    paragraphs = [b.text for b in doc.blocks if b.kind == "paragraph"]

    assert len(paragraphs) == doc.fields["messages"]
    assert paragraphs[0].startswith("2025-06-03 08:55 ")


def test_docx_procedure_table_matches_source_steps() -> None:
    rows = procedure_rows(parse("greenfield/sop-db-001.docx"))

    assert [row[1] for row in rows] == source_steps("sop-db-001")


def test_docx_skips_struck_through_text(tmp_path: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Restart the workers. ")
    paragraph.add_run("Then delete the queue.").font.strike = True
    path = tmp_path / "struck.docx"
    document.save(str(path))

    doc = parse_file(path)

    assert "Restart the workers." in doc.text
    assert "delete the queue" not in doc.text


def test_unstyled_docx_reads_steps_as_list_items() -> None:
    doc = parse("brownfield/lag-fix-old.docx")

    assert [b.text[:2] for b in doc.blocks if b.kind == "list_item"] == ["1)", "2)", "3)"]


def test_pdf_drops_running_headers_and_page_numbers() -> None:
    doc = parse("greenfield/sop-queue-001.pdf")

    assert "Page 1 of" not in doc.text
    assert doc.text.count("Document number") == 1


def test_pdf_procedure_table_matches_source_steps() -> None:
    rows = procedure_rows(parse("greenfield/sop-queue-001.pdf"))

    assert [row[1] for row in rows] == source_steps("sop-queue-001")


@pytest.mark.parametrize("path", sorted(CORPUS.glob("*/*.pdf")), ids=lambda path: path.name)
def test_pdf_text_has_no_unmapped_glyphs(path: Path) -> None:
    assert "(cid:" not in parse_file(path).text


@pytest.mark.parametrize("relative", ["greenfield/sop-db-001.docx", "greenfield/sop-cache-001.md"])
def test_table_of_contents_is_dropped(relative: str) -> None:
    doc = parse(relative)

    assert doc.text.count("Appendix A. Revision history") == 1
    assert all(b.text != "Table of contents" for b in doc.blocks)


def test_identical_files_get_identical_hashes() -> None:
    copy = CORPUS / "brownfield" / "cache-down-runbook-copy.md"
    original = CORPUS / "greenfield" / "sop-cache-001.md"

    assert raw_sha256(copy.read_bytes()) == raw_sha256(original.read_bytes())
    assert content_hash(parse_file(copy)) == content_hash(parse_file(original))


def test_store_raw_writes_each_file_once(tmp_path: Path) -> None:
    first = store_raw(b"runbook", tmp_path)
    second = store_raw(b"runbook", tmp_path)

    assert first == second
    assert [path.name for path in tmp_path.iterdir()] == [first]


def test_unknown_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedFormat):
        parse_bytes(b"data", "notes.xyz")

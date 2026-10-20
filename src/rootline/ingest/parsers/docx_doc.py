import io
from collections.abc import Iterator

import docx
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

from rootline.core.models import ParsedDoc
from rootline.ingest.document_builder import DocBuilder
from rootline.ingest.parsers.text_doc import add_text

STRUCTURE_STYLES = ("Title", "Heading", "List")


def parse(data: bytes, source_key: str) -> ParsedDoc:
    document = docx.Document(io.BytesIO(data))
    items = list(document.iter_inner_content())
    builder = DocBuilder()
    if any(isinstance(i, Paragraph) and style_of(i).startswith(STRUCTURE_STYLES) for i in items):
        add_styled(builder, items)
    else:
        add_unstyled(builder, items)
    properties = document.core_properties
    fields = {"core_title": properties.title, "core_author": properties.author}
    return builder.build(source_key, "docx", {k: v for k, v in fields.items() if v})


def add_styled(builder: DocBuilder, items: list[Paragraph | Table]) -> None:
    for item in items:
        if isinstance(item, Table):
            builder.add_table(table_rows(item))
            continue
        style, text = style_of(item), visible_text([item])
        if style == "Title":
            builder.add("heading", text, level=1)
        elif style.startswith("Heading"):
            level = style.split()[-1]
            builder.add("heading", text, level=int(level) if level.isdigit() else 1)
        elif style.startswith("List"):
            builder.add("list_item", text)
        else:
            builder.add("paragraph", text)


def add_unstyled(builder: DocBuilder, items: list[Paragraph | Table]) -> None:
    # Without Word styles the only structure is in the text itself, as in a plain text file.
    lines: list[str] = []
    title_pending = True
    for item in items:
        if isinstance(item, Paragraph):
            lines.append(visible_text([item]))
            continue
        title_pending = add_lines_as_text(builder, lines, title_pending)
        lines = []
        builder.add_table(table_rows(item))
    add_lines_as_text(builder, lines, title_pending)


def add_lines_as_text(builder: DocBuilder, lines: list[str], title_pending: bool) -> bool:
    text = "\n".join(lines)
    if not text.strip():
        return title_pending
    add_text(builder, text, first_line_is_title=title_pending)
    return False


def table_rows(table: Table) -> list[list[str]]:
    return [[visible_text(cell.paragraphs) for cell in row.cells] for row in table.rows]


def style_of(paragraph: Paragraph) -> str:
    style = paragraph.style
    return style.name if style is not None and style.name else ""


def visible_text(paragraphs: list[Paragraph]) -> str:
    return "\n".join("".join(run_texts(paragraph)) for paragraph in paragraphs)


def run_texts(paragraph: Paragraph) -> Iterator[str]:
    # Struck-through text is deleted content in many legacy docs, such as retired steps.
    for item in paragraph.iter_inner_content():
        for run in item.runs if isinstance(item, Hyperlink) else [item]:
            if not run.font.strike:
                yield run.text

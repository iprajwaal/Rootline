import io
import zipfile
from datetime import datetime
from pathlib import Path
from typing import assert_never

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph as DocxParagraph

from rootline.corpus.layout import COMPANY, Bullets, Element, Heading, Layout, Paragraph, Table


def write_docx(layout: Layout, path: Path) -> None:
    document = Document()
    add_page_header(document, layout.header_lines)
    add_page_footer(document)
    for element in layout.elements:
        add_element(document, element)

    stamp = datetime.fromisoformat(layout.date) if layout.date else datetime(1980, 1, 1)
    path.write_bytes(docx_bytes(document, layout.title, layout.author, stamp))


def docx_bytes(document: DocxDocument, title: str, author: str, stamp: datetime) -> bytes:
    properties = document.core_properties
    properties.title = title
    properties.author = author
    properties.last_modified_by = author
    properties.created = properties.modified = stamp

    buffer = io.BytesIO()
    document.save(buffer)
    return with_fixed_zip_times(buffer.getvalue(), stamp)


def with_fixed_zip_times(data: bytes, stamp: datetime) -> bytes:
    # python-docx stamps every zip entry with the current time, so the same source would
    # produce a different file on each render. Pin the times to keep renders identical.
    fixed = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as original,
        zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in original.infolist():
            info.date_time = (stamp.year, stamp.month, stamp.day, 0, 0, 0)
            target.writestr(info, original.read(info.filename))
    return fixed.getvalue()


def add_element(document: DocxDocument, element: Element) -> None:
    match element:
        case Heading(text, level):
            document.add_heading(text, level=level)
        case Paragraph(text):
            document.add_paragraph(text)
        case Bullets(items, numbered):
            # Plain "1." prefixes: Word list styles keep counting across separate lists.
            for n, item in enumerate(items, start=1):
                if numbered:
                    document.add_paragraph(f"{n}. {item}")
                else:
                    document.add_paragraph(item, style="List Bullet")
        case Table(header, rows, code_columns):
            add_table(document, header, rows, code_columns)
        case _:
            assert_never(element)


def add_table(
    document: DocxDocument, header: list[str], rows: list[list[str]], code_columns: frozenset[int]
) -> None:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        for index, (cell, text) in enumerate(zip(table.add_row().cells, row, strict=True)):
            run = cell.paragraphs[0].add_run(text)
            if index in code_columns:
                run.font.name = "Consolas"
    document.add_paragraph()


def add_page_header(document: DocxDocument, lines: list[str]) -> None:
    header = document.sections[0].header
    company = header.paragraphs[0]
    company.add_run(COMPANY).bold = True
    for line in lines:
        paragraph = header.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_page_footer(document: DocxDocument) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.add_run("Internal    Page ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" of ")
    add_field(paragraph, "NUMPAGES")
    for run in paragraph.runs:
        run.font.size = Pt(8)


def add_field(paragraph: DocxParagraph, code: str) -> None:
    # python-docx has no API for fields, so the page number is written as raw field XML.
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])

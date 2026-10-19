import html
import io
import json
import re
import textwrap
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from hashlib import sha256
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import docx
import yaml
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Flowable, SimpleDocTemplate, Spacer
from reportlab.platypus import Paragraph as PdfParagraph

from rootline.corpus.docx_writer import docx_bytes
from rootline.corpus.layout import EMPTY, Source, layout_for, load_document
from rootline.corpus.md_writer import to_markdown
from rootline.models import DocMeta, RcaRecord, SopRecord

CORPUS = Path("corpus")
PLAN_FILE = "sources/brownfield_plan.yaml"
MAIL_DOMAIN = "shoplite.example"

# Sections a rushed author leaves empty in a half-filled template.
SOP_BLANKS: dict[str, Any] = {
    "purpose": "",
    "scope": "",
    "definitions": [],
    "roles": [],
    "related_documents": [],
    "troubleshooting": [],
    "compliance": [],
    "review_cycle": "",
    "records": "",
    "revision_history": [],
}
RCA_BLANKS: dict[str, Any] = {
    "authors": [],
    "stakeholders": [],
    "impact_description": "",
    "impact_metrics": "",
    "customer_feedback": "",
    "five_whys": [],
    "fishbone": [],
    "action_items": [],
}

WIKI_PAGE = """<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><title>Orders Engineering : {title}</title></head>
<body>
<div id="header"><a href="/wiki">ShopLite Wiki</a> | <a href="/wiki/search">Search</a> \
| <a href="/wiki/login">Log in</a></div>
<div id="breadcrumbs">Pages / Orders Engineering / {section}</div>
<h1 id="title-heading">{title}</h1>
<div class="page-metadata">Created by {author}, last modified on {updated}</div>
<div id="main-content" class="wiki-content">
{content}
</div>
<div id="footer">Exported from the wiki. This page may be out of date.</div>
</body>
</html>
"""


@dataclass(frozen=True)
class Doc:
    meta: DocMeta
    record: SopRecord | RcaRecord


@dataclass(frozen=True)
class Spec:
    file: str
    transform: str
    sources: list[str]
    canaries: list[str] = field(default_factory=list)
    title: str | None = None
    misleading: bool = False
    outdated: dict[str, str] = field(default_factory=dict)
    copy_of: str | None = None


Transform = Callable[[Spec, list[Doc], list[str]], str | bytes]


def generate(corpus_dir: Path, output_dir: Path, truth_path: Path) -> list[dict[str, Any]]:
    plan: dict[str, Any] = yaml.safe_load((corpus_dir / PLAN_FILE).read_text())
    greenfield = load_docs(corpus_dir / "sources" / "greenfield")
    docs = greenfield | load_docs(corpus_dir / "sources" / "brownfield")
    output_dir.mkdir(parents=True, exist_ok=True)
    truth_path.parent.mkdir(parents=True, exist_ok=True)

    rows = []
    # Seeded with the greenfield files so a brownfield copy of one is recorded too.
    first_file_by_hash = {
        sha256(path.read_bytes()).hexdigest(): f"greenfield/{path.name}"
        for path in sorted((corpus_dir / "greenfield").iterdir())
    }
    for entry in plan["documents"]:
        spec = Spec(**entry)
        sources = [with_replacements(docs[doc_id], spec.outdated) for doc_id in spec.sources]
        canaries = [plan["canaries"][name] for name in spec.canaries]
        if spec.copy_of:
            data = (corpus_dir / spec.copy_of).read_bytes()
        else:
            content = TRANSFORMS[spec.transform](spec, sources, [c["line"] for c in canaries])
            data = content.encode() if isinstance(content, str) else content
        (output_dir / spec.file).write_bytes(data)

        digest = sha256(data).hexdigest()
        row = truth_row(spec, sources, set(greenfield), [str(c["value"]) for c in canaries])
        row["same_bytes_as"] = first_file_by_hash.get(digest)
        first_file_by_hash.setdefault(digest, row["file"])
        rows.append(row)

    truth_path.write_text("".join(json.dumps(row) + "\n" for row in rows))
    return rows


def load_docs(source_dir: Path) -> dict[str, Doc]:
    docs = {}
    for path in sorted(source_dir.glob("*.yaml")):
        meta, record = load_document(path)
        docs[meta.doc_id] = Doc(meta, record)
    return docs


def with_replacements(doc: Doc, replacements: dict[str, str]) -> Doc:
    meta_json = doc.meta.model_dump_json()
    record_json = doc.record.model_dump_json()
    for old, new in replacements.items():
        meta_json = meta_json.replace(old, new)
        record_json = record_json.replace(old, new)
    return Doc(
        DocMeta.model_validate_json(meta_json), type(doc.record).model_validate_json(record_json)
    )


def truth_row(
    spec: Spec, docs: list[Doc], greenfield_ids: set[str], canary_values: list[str]
) -> dict[str, Any]:
    return {
        "file": f"brownfield/{spec.file}",
        "format": Path(spec.file).suffix.lstrip("."),
        "transform": spec.transform,
        "doc_type": "mixed" if len(docs) > 1 else docs[0].meta.doc_type,
        "source_ids": spec.sources,
        "services": sorted({service for doc in docs for service in doc.meta.services}),
        "greenfield_twins": [doc_id for doc_id in spec.sources if doc_id in greenfield_ids],
        "outdated": bool(spec.outdated),
        "stale_values": list(spec.outdated.values()),
        "misleading_title": spec.misleading,
        "canaries": canary_values,
    }


def plain_runbook(spec: Spec, docs: list[Doc], extras: list[str]) -> str:
    return "\n".join(runbook_lines(spec, docs, extras)) + "\n"


def old_word(spec: Spec, docs: list[Doc], extras: list[str]) -> bytes:
    title, *body = runbook_lines(spec, docs, extras)
    document = docx.Document()
    document.add_paragraph().add_run(title).bold = True
    for line in body:
        document.add_paragraph(line)
    meta = docs[0].meta
    return docx_bytes(document, title, meta.owner, day_of(meta.effective_date))


def pdf_export(spec: Spec, docs: list[Doc], extras: list[str]) -> bytes:
    lines = runbook_lines(spec, docs, extras)
    style = getSampleStyleSheet()["BodyText"]
    flowables: list[Flowable] = [
        PdfParagraph(escape(line), style) if line else Spacer(1, 6) for line in lines
    ]
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, title=lines[0], invariant=True).build(flowables)
    return buffer.getvalue()


def runbook_lines(spec: Spec, docs: list[Doc], extras: list[str]) -> list[str]:
    meta, sop = docs[0].meta, sop_of(docs[0])
    lines = [
        spec.title or meta.title.lower(),
        f"(copied from the old wiki, owner: {meta.owner.lower()}, "
        f"last checked {meta.effective_date or 'unknown'})",
        "",
        sop.when_to_use,
    ]
    if sop.prerequisites:
        lines += ["", "before you start: " + " ".join(sop.prerequisites)]
    lines += ["", "steps"]
    for step in sop.steps:
        lines.append(f"{step.ordinal}) {step.instruction}")
        if step.command:
            lines.append(f"   {step.command}")
        if step.expect:
            lines.append(f"   should see: {step.expect}")
    if sop.rollback:
        lines += ["", "to undo"]
        lines += [f"- {step.instruction} {step.command or ''}".rstrip() for step in sop.rollback]
    if extras:
        lines += ["", "notes", *extras]
    if sop.escalation:
        lines += ["", f"if this does not help: {lower_first(sop.escalation[0])}"]
    return lines


def email_thread(spec: Spec, docs: list[Doc], extras: list[str]) -> str:
    meta, rca = docs[0].meta, rca_of(docs[0])
    sender = (rca.authors or [meta.owner])[0]
    replied_to = (rca.stakeholders[1:] or [meta.owner])[0]
    sent = day_of(rca.incident_date or meta.effective_date) + timedelta(days=1, hours=10)
    lines = [
        "Thanks all. Writing down what we know so far, proper RCA to follow.",
        "",
        f"What happened: {rca.short_description}",
        f"Impact: {rca.impact_metrics}",
        "",
        "Timeline (UTC):",
        *[f"{e.timestamp} {e.event} {e.action}" for e in rca.timeline],
        "",
        f"Root cause: {' '.join(rca.root_causes)}",
        f"Fix: {rca.immediate_resolution}",
        "",
        "Follow ups:",
        *[f"- {item.action} ({item.owner})" for item in rca.action_items],
    ]
    if extras:
        lines += ["", *extras]
    lines += [
        "",
        "--",
        sender,
        meta.owner,
        "",
        f"On {sent - timedelta(hours=10):%a, %d %b %Y}, {replied_to} wrote:",
        f"> {rca.trigger} fired, is anyone looking?",
        *[f"> {line}" for line in textwrap.wrap(rca.problem_statement, 70)],
    ]

    message = EmailMessage()
    message["From"] = f"{sender} <{slug(sender, '.')}@{MAIL_DOMAIN}>"
    message["To"] = f"orders-eng@{MAIL_DOMAIN}"
    message["Subject"] = spec.title or f"RE: RE: {meta.title.lower()}"
    message["Date"] = format_datetime(sent)
    message["Message-ID"] = f"<{sha256(spec.file.encode()).hexdigest()[:16]}@{MAIL_DOMAIN}>"
    message.set_content("\n".join(lines) + "\n")
    return message.as_string()


def slack_export(spec: Spec, docs: list[Doc], extras: list[str]) -> str:
    meta, rca = docs[0].meta, rca_of(docs[0])
    messages = [
        slack_message(e.responsibility, parse_utc(e.timestamp), f"{e.event} {e.action}")
        for e in rca.timeline
    ]
    last = parse_utc(rca.timeline[-1].timestamp)
    for minutes, line in enumerate(extras, start=1):
        messages.append(slack_message("SRE on-call", last + timedelta(minutes=minutes), line))
    summary = f"summary: {' '.join(rca.root_causes)} fix: {rca.immediate_resolution}"
    messages.append(slack_message(meta.owner, last + timedelta(minutes=30), summary))
    channel = spec.title or f"inc-{rca.incident_date}-{slug(meta.title)}"
    return json.dumps({"channel": channel, "messages": messages}, indent=2) + "\n"


def slack_message(who: str, at: datetime, text: str) -> dict[str, str]:
    return {
        "type": "message",
        "user_name": slug(who),
        "ts": f"{int(at.timestamp())}.000100",
        "text": text,
    }


def wiki_html(spec: Spec, docs: list[Doc], extras: list[str]) -> str:
    first = docs[0]
    parts: list[str] = []
    for doc in docs:
        parts += sop_html(doc.record) if isinstance(doc.record, SopRecord) else rca_html(doc.record)
    parts += [f"<p>{html.escape(line)}</p>" for line in extras]
    author = first.record.authors[0] if isinstance(first.record, RcaRecord) else first.meta.owner
    return WIKI_PAGE.format(
        title=html.escape(spec.title or first.meta.title),
        section="Runbooks" if first.meta.doc_type == "sop" else "Incidents",
        author=html.escape(author),
        updated=first.meta.effective_date or "unknown",
        content="\n".join(parts),
    )


def sop_html(sop: SopRecord) -> list[str]:
    e = html.escape
    parts = [f"<p>{e(sop.purpose)}</p>"] if sop.purpose else []
    parts += [
        "<h2>When to use</h2>",
        f"<p>{e(sop.when_to_use)}</p>",
        "<h2>Steps</h2>",
        "<table>",
        "<tr><th>#</th><th>What to do</th><th>Command</th><th>Expected</th></tr>",
        *[
            f"<tr><td>{s.ordinal}</td><td>{e(s.instruction)}</td>"
            f"<td><code>{e(s.command or '')}</code></td><td>{e(s.expect or '')}</td></tr>"
            for s in sop.steps
        ],
        "</table>",
    ]
    if sop.rollback:
        parts += ["<h2>Rollback</h2>", "<ol>"]
        parts += [
            f"<li>{e(s.instruction)} <code>{e(s.command or '')}</code></li>" for s in sop.rollback
        ]
        parts.append("</ol>")
    if sop.escalation:
        parts += [
            "<h2>Contacts</h2>",
            "<ul>",
            *[f"<li>{e(x)}</li>" for x in sop.escalation],
            "</ul>",
        ]
    return parts


def rca_html(rca: RcaRecord) -> list[str]:
    e = html.escape
    return [
        "<h2>Summary</h2>",
        f"<p>{e(rca.short_description)}</p>",
        "<h2>What happened</h2>",
        f"<p>{e(rca.problem_statement)}</p>",
        "<h2>Timeline</h2>",
        "<table>",
        "<tr><th>Time (UTC)</th><th>What</th><th>Who</th></tr>",
        *[
            f"<tr><td>{e(t.timestamp)}</td><td>{e(t.event)} {e(t.action)}</td>"
            f"<td>{e(t.responsibility)}</td></tr>"
            for t in rca.timeline
        ],
        "</table>",
        "<h2>Why it happened</h2>",
        "<ul>",
        *[f"<li>{e(cause)}</li>" for cause in rca.root_causes],
        "</ul>",
        "<h2>Follow-ups</h2>",
        "<ul>",
        *[f"<li>{e(item.action)} ({e(item.owner)})</li>" for item in rca.action_items],
        "</ul>",
    ]


def half_filled(spec: Spec, docs: list[Doc], extras: list[str]) -> str:
    doc = docs[0]
    blanks = SOP_BLANKS if isinstance(doc.record, SopRecord) else RCA_BLANKS
    record = doc.record.model_copy(update=blanks)
    meta = doc.meta.model_copy(
        update={"doc_id": "TBD", "related": [], "title": spec.title or doc.meta.title}
    )
    text = to_markdown(layout_for(Source("md", meta, record))).replace(EMPTY, "TBD")
    if extras:
        text += "\nNotes:\n" + "".join(f"- {line}\n" for line in extras)
    return text


def sop_of(doc: Doc) -> SopRecord:
    if not isinstance(doc.record, SopRecord):
        raise ValueError(f"{doc.meta.doc_id} is not an SOP")
    return doc.record


def rca_of(doc: Doc) -> RcaRecord:
    if not isinstance(doc.record, RcaRecord):
        raise ValueError(f"{doc.meta.doc_id} is not an RCA")
    return doc.record


def day_of(date: str | None) -> datetime:
    return datetime.fromisoformat(date or "1980-01-01").replace(tzinfo=UTC)


def parse_utc(timestamp: str) -> datetime:
    return datetime.strptime(timestamp, "%Y-%m-%d %H:%M").replace(tzinfo=UTC)


def slug(text: str, separator: str = "-") -> str:
    return re.sub(r"[^a-z0-9]+", separator, text.lower()).strip(separator)[:40]


def lower_first(text: str) -> str:
    return text[:1].lower() + text[1:]


TRANSFORMS: dict[str, Transform] = {
    "plain_runbook": plain_runbook,
    "old_word": old_word,
    "pdf_export": pdf_export,
    "email_thread": email_thread,
    "slack_export": slack_export,
    "wiki_html": wiki_html,
    "half_filled": half_filled,
}


def main() -> None:
    rows = generate(CORPUS, CORPUS / "brownfield", CORPUS / "ground_truth" / "brownfield.jsonl")
    print(f"brownfield files={len(rows)} dir={CORPUS / 'brownfield'}")


if __name__ == "__main__":
    main()
